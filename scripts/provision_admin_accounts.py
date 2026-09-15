"""Read the supplied workbook into a private, hash-only administrator manifest."""
import argparse
import json
import os
from pathlib import Path
import sys
import tempfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "apps/api"))
from app.auth_security import hash_secret, normalize_login, verify_secret


def provision(source: Path, destination: Path) -> int:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = [str(value or "").strip() for value in rows[0]] if rows else []
    required = ["사용자명", "계정(아이디)", "비밀번호"]
    if any(key not in headers for key in required):
        raise ValueError("Required account columns are missing")
    previous = {}
    if destination.exists():
        previous = {entry["login"]: entry["password_hash"] for entry in json.loads(destination.read_text())["accounts"]}
    accounts, seen = [], set()
    for row in rows[1:]:
        if not any(value is not None for value in row):
            continue
        name, login, password = (row[headers.index(key)] for key in required)
        if not all(isinstance(value, str) and value.strip() for value in [name, login, password]):
            raise ValueError("Account cells must contain non-empty text")
        login, name = normalize_login(login), name.strip()
        if len(login) > 80 or len(name) > 80 or len(password) > 200 or login in seen:
            raise ValueError("Invalid or duplicate account")
        seen.add(login)
        encoded = previous.get(login, "")
        if not verify_secret(password, encoded):
            encoded = hash_secret(password)
        accounts.append({"login": login, "name": name, "password_hash": encoded})
    if not accounts:
        raise ValueError("At least one administrator is required")
    destination.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    fd, temporary = tempfile.mkstemp(dir=destination.parent, prefix=".accounts-")
    try:
        with os.fdopen(fd, "w") as output:
            json.dump({"version": 1, "accounts": accounts}, output, ensure_ascii=False, indent=2)
            output.write("\n")
        os.chmod(temporary, 0o600)
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return len(accounts)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "runtime/auth/admin-accounts.json")
    args = parser.parse_args()
    try:
        count = provision(args.workbook, args.output)
    except Exception:
        # Do not let parser errors or workbook values reach shell logs.
        raise SystemExit("Account import failed. Check workbook columns, values and destination permissions.") from None
    print(f"Prepared {count} administrator accounts in a private hash manifest.")
